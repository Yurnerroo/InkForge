from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE

from inkforge.content_checker.docreader import read_text_metrics


def test_read_txt_single_line_is_title_only(tmp_path: Path) -> None:
    path = tmp_path / "article.txt"
    path.write_text("Одне два три чотири.", encoding="utf-8")

    result = read_text_metrics(path)

    assert result.title == "Одне два три чотири."
    assert result.lead == ""
    assert result.body == ""
    assert result.word_count == 4
    assert result.warnings == []


def test_read_txt_two_lines_is_title_and_body_no_lead(tmp_path: Path) -> None:
    path = tmp_path / "article.txt"
    path.write_text("Короткий заголовок.\nДовший основний текст статті з деталями.", encoding="utf-8")

    result = read_text_metrics(path)

    assert result.title == "Короткий заголовок."
    assert result.lead == ""
    assert result.body == "Довший основний текст статті з деталями."
    assert result.has_lead_paragraph is False


def test_read_txt_three_lines_detects_short_second_line_as_lead(tmp_path: Path) -> None:
    path = tmp_path / "article.txt"
    path.write_text(
        "Заголовок статті\nКороткий лід.\nДовший основний текст статті з деталями.",
        encoding="utf-8",
    )

    result = read_text_metrics(path)

    assert result.title == "Заголовок статті"
    assert result.lead == "Короткий лід."
    assert result.body == "Довший основний текст статті з деталями."
    assert result.has_lead_paragraph is True


def test_read_txt_long_second_line_goes_to_body_not_lead(tmp_path: Path) -> None:
    long_line = "Д" * 300
    path = tmp_path / "article.txt"
    path.write_text(f"Заголовок\n{long_line}\nТретій рядок тіла.", encoding="utf-8")

    result = read_text_metrics(path)

    assert result.title == "Заголовок"
    assert result.lead == ""
    assert result.body == f"{long_line}\nТретій рядок тіла."


def test_read_docx_detects_title_and_lead_styles(tmp_path: Path) -> None:
    document = Document()
    document.styles.add_style("Заголовок", WD_STYLE_TYPE.PARAGRAPH)
    document.styles.add_style("Vrizka", WD_STYLE_TYPE.PARAGRAPH)

    title = document.add_paragraph("Гучний заголовок")
    title.style = document.styles["Заголовок"]
    lead = document.add_paragraph("Короткий лід-абзац на початку.")
    lead.style = document.styles["Vrizka"]
    document.add_paragraph("Основний текст статті з кількома словами тут.")
    document.add_paragraph("Другий абзац тіла статті.")

    path = tmp_path / "article.docx"
    document.save(str(path))

    result = read_text_metrics(path)

    assert result.title == "Гучний заголовок"
    assert result.lead == "Короткий лід-абзац на початку."
    assert result.body == (
        "Основний текст статті з кількома словами тут.\nДругий абзац тіла статті."
    )
    assert result.has_lead_paragraph is True
    assert result.word_count > 0
    assert result.warnings == []


def test_read_docx_lead_style_without_title_style(tmp_path: Path) -> None:
    document = Document()
    document.styles.add_style("Vrizka", WD_STYLE_TYPE.PARAGRAPH)
    lead = document.add_paragraph("Короткий лід-абзац на початку.")
    lead.style = document.styles["Vrizka"]
    document.add_paragraph("Основний текст статті з кількома словами тут.")
    path = tmp_path / "article.docx"
    document.save(str(path))

    result = read_text_metrics(path)

    assert result.title == ""
    assert result.lead == "Короткий лід-абзац на початку."
    assert result.has_lead_paragraph is True
    assert result.warnings == []


def test_read_docx_without_recognized_styles_uses_positional_fallback(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("Звичайний перший абзац без спеціального стилю.")
    document.add_paragraph("Другий абзац, теж без стилю, довший за лід-поріг тому що тіло статті.")
    path = tmp_path / "article.docx"
    document.save(str(path))

    result = read_text_metrics(path)

    assert result.title == "Звичайний перший абзац без спеціального стилю."
    assert result.has_lead_paragraph is False
    assert any("positional fallback" in w for w in result.warnings)


def test_read_text_metrics_missing_file(tmp_path: Path) -> None:
    result = read_text_metrics(tmp_path / "missing.txt")

    assert result.word_count == 0
    assert any("not found" in w for w in result.warnings)


def test_read_text_metrics_unsupported_legacy_doc(tmp_path: Path) -> None:
    path = tmp_path / "article.doc"
    path.write_text("legacy", encoding="utf-8")

    result = read_text_metrics(path)

    assert any("not supported" in w for w in result.warnings)


def test_read_text_metrics_unknown_extension(tmp_path: Path) -> None:
    path = tmp_path / "article.rtf"
    path.write_text("data", encoding="utf-8")

    result = read_text_metrics(path)

    assert any("Unsupported text format" in w for w in result.warnings)
